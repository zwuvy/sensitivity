'''
Sensitivity Test by Karan Gill

READ THE INSTRUCTIONS DOCUMENT BEFORE RUNNING THIS PROGRAM

This program will perform 7 total sensitivity tests. The user provides a shift and step for the bit rate to change at.
The program will then do 3 test slower, 3 tests faster, and 1 test at the default rate using this information.
1000 signals will be sent at each power level (see power_level_list below) at 250ms per signal.

I have only seen this program bug out at the beginning when setting the frequency or changing windows, 
for some reason it will just miss it.
Therefore I recommend manually setting the frequency on the machine before running the program and also watching the program
for about 1 minute to make sure it switches from the generator window, back to the vscode window, and then again back to the 
generator window to begin testing.
Another common issue is minimizing the vscode or microsoft edge window with the VNC.
Make sure these windows are not minimized so that the program can bring them to the foreground.

The results will be automatically put into an excel file titled "results".
This excel file will be overwritten everytime the program is run so copy the data to another location.
This excel file will be in the location that the terminal is currently in (see the bottom window in vscode)
For me, this was my user folder
C Drive -> Users -> jgillk1

'''

import pyautogui
import win32gui
import time
import pyperclip
import openpyxl 

#invisible characters in these strings
generator_window = "noVNC - Work - Microsoft​ Edge"  
vscode_window = "alpha.py - Visual Studio Code"

power_level_list = [-90, -100, -105, -110, -111, -112, -113, -114, -115, -116, -117, -118, -119, -120]

def init_baud_symbol_input():
    baud_rate = input("Enter '1' if this is a 1K baud device and '2' if it is a 2K baud device\n")
    while baud_rate != "1" and baud_rate != "2":
        print("Invalid input: enter either '1' or '2'")
        baud_rate = input()

    if baud_rate == "1":
        return 1886.79245
    if baud_rate == "2":
        return 4000.00000

def init_frequency_input():
    frequency_level = input("Enter '433' if this is a 433.92 MHz device and '868' if it is a 868.35 MHz device\n")
    while frequency_level != "433" and frequency_level != "868":
        print("Invalid input: enter either '433' or '868'")
        frequency_level = input()

    if frequency_level == "433":
        return 433.92
    if frequency_level == "868":
        return 868.35

def init_baud_symbol_percent_shift_input():
    percent_shift = input("Enter the starting percent shift from the default symbol rate (0% to 4%)\n")
    percent_shift = percent_shift.strip("%")
    percent_shift = float(percent_shift)

    while percent_shift < 0.0 or percent_shift > 4.00:
        print("Invalid input: enter a shift of 0.0 or greater and 4.00 or less")
        percent_shift = input()
        percent_shift = percent_shift.strip("%")
        percent_shift = float(percent_shift)

    return percent_shift/100.0

def init_baud_symbol_percent_step_input():
    percent_step = input("Enter the percent step to increase/decrease from the percent shift (0% to 2%)\n")
    percent_step = percent_step.strip("%")
    percent_step = float(percent_step)

    while percent_step < 0.0 or percent_step > 2.0:
        print("Invalid input: enter a step of 0.0 or greater and 2.0 or less")
        percent_step = input()
        percent_step = percent_step.strip("%")
        percent_step = float(percent_step)
    
    return percent_step/100.0

def init_workbook(rate_list, percent_list):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws['A1'] = "RX Power"
    ws['B1'] = "Symbol Rate"
    ws['B2'] = "Percent Shift"
    num = 3
    for val in power_level_list:
        ws["A" + str(num)] = val
        num += 1
    
    cell = "C"
    for i in range(len(rate_list)):
        ws[cell + "1"] = str(rate_list[i])
        ws[cell + "2"] = str(percent_list[i] * 100) + "%"
        cell = chr(ord(cell) + 1)
    
    wb.save("results.xlsx")
    wb.close()

def symbol_rate_conversion(shift, step, symbol_rate):
    percent_shift_list = [-shift - 2*step, -shift - step, -shift, 0, shift, shift + step, shift + 2*step]
    symbol_rate_list = []
    for value in percent_shift_list:
        symbol_rate_list.append(round(value*symbol_rate + symbol_rate, 3))

    return symbol_rate_list, percent_shift_list

def switch_window(window_title):
    time.sleep(1)
    hwnd = win32gui.FindWindow(None, window_title)
    win32gui.SetForegroundWindow(hwnd)
    time.sleep(1)

def go_home():
    time.sleep(1)
    pyautogui.click(x = 1550, y = 250) #return to main screen
    time.sleep(0.25)
    pyautogui.click(x = 1550, y = 250)
    time.sleep(1)

def half_second_click(x_coord, y_coord):
    pyautogui.click(x = x_coord, y = y_coord)
    time.sleep(0.5)

def set_frequency(frequency):
    switch_window(generator_window)
    switch_window(generator_window)
    go_home()
    half_second_click(1550, 230)
    half_second_click(500, 300)
    pyautogui.typewrite(str(frequency) + "\n")

def set_power_level(power_level):
    switch_window(generator_window)
    go_home()
    half_second_click(1400, 300)
    pyautogui.typewrite(str(power_level) + "\n")

def set_symbol_rate(symbol_rate):
    switch_window(generator_window)
    go_home()
    half_second_click(500, 550)
    half_second_click(400, 250)
    half_second_click(750, 550)
    half_second_click(1100, 550)
    pyautogui.typewrite(str(symbol_rate) + "\n")
    go_home()

def transmit_1000_signals():
    switch_window(generator_window)
    go_home()
    half_second_click(500, 550)

    #Change the range to alter the number of signals sent per test
    #Change the number in sleep to change the time between sending signals (in seconds)
    for i in range(1000):
        pyautogui.click(750, 450)
        time.sleep(0.25)

    go_home()

def copy_from_serial_monitor():
    switch_window(vscode_window)
    pyautogui.doubleClick(1700, 900)
    time.sleep(0.5)
    pyautogui.hotkey("ctrl", "c")
    time.sleep(0.5)

def paste_to_excel(column, row, total_recieved):
    time.sleep(0.5)
    wb = openpyxl.load_workbook("results.xlsx")
    ws = wb.active
    new_recieved = int(pyperclip.paste()) - total_recieved
    ws[column + str(row)] = new_recieved
    wb.save("results.xlsx")
    wb.close()

def main():
    start_time = time.time()
    
    symbol_rate = init_baud_symbol_input()
    frequency = init_frequency_input()
    symbol_rate_shift_percent = init_baud_symbol_percent_shift_input()
    symbol_rate_step_percent = init_baud_symbol_percent_step_input()
    rate_list, percent_list = symbol_rate_conversion(symbol_rate_shift_percent, symbol_rate_step_percent, symbol_rate)
    init_workbook(rate_list, percent_list)
    

    switch_window(generator_window)
    switch_window(vscode_window)


    #frequency only needs to be set once, tends to bug, recommend manually setting frequency before testing
    set_frequency(frequency)

    column = "C"
    row = 3
    total_recieved = 0

    for rate in rate_list:
        set_symbol_rate(rate)
        print("Currently Testing: ", rate, "sym / s")
        for power in power_level_list:
            print("@", power, "dBm")
            set_power_level(power)
            transmit_1000_signals()
            copy_from_serial_monitor()
            paste_to_excel(column, row, total_recieved)
            total_recieved = int(pyperclip.paste())
            row += 1
        if symbol_rate_shift_percent == 0.0:
            break

        row = 3
        column = chr(ord(column) + 1)

    switch_window(vscode_window)
    end_time = time.time()
    total_seconds = end_time - start_time
    total_minutes = total_seconds / 60.0
    total_hours = total_minutes / 60.0
    print("ALL DONE")
    print("ALL DONE")
    print("ALL DONE")
    print("Testing completed in ", round(total_minutes, 3), "minutes /", round(total_hours, 3), "hours")
    print("The excel file's location is in the directory that the terminal is in, likely: C Drive -> Users -> Your ID")
    
main()